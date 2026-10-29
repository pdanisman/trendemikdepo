#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TrEndemik - Test planı: fotoğrafları seç, indir, test şablonunu doldur
=======================================================================
Girdi : endemik_veri_seti_v2.xlsx  +  kontrol_veri_seti.xlsx  +  test_sablonu.xlsx (boş şablon)
Çıktı : fotograflar/  (F0001.jpg ...)  +  test_sablonu_dolu.xlsx  (Fotograf_Kaydi doldurulmuş)

Seçim kuralı (önceden yazılmış, sonuçlara bakılarak DEĞİŞTİRİLMEZ):
  * Endemik (P1): "Test_İçin_Uygun = Evet" olan HER türden 1 fotoğraf   -> türler arası bağımsızlık
  * Kontrol (N1/N2/N3): kontrol_veri_seti.xlsx'teki HER türden 1 fotoğraf
  * Aday sırası: iNaturalist (araştırma düzeyi) > GBIF > Commons; aynı kaynak içinde sabit tohumlu rastgele.
    İlk adayın indirilemezse/bozuksa sıradaki denenir; hiçbiri indirilemezse tür dışarıda kalır (raporlanır).
  * Fotoğraf numaraları (F0001...) tüm türler tohumlu karıştırılarak verilir; numaradan grup anlaşılmaz.
  * GELISTIRME / TEST ayrımı: her grup (P1 için yayılış grubu da) içinde ayrı ayrı, %30 / %70.

Aşamalar:  indir  (seç + indir)   |   plan  (şablonu doldur)   |   hepsi
Kullanım:
  python test_plani_olustur.py --asama hepsi
"""
import argparse
import json
import os
import random
import time

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import MultiCellRange

HEADERS = {"User-Agent": "TrEndemikArastirma/1.0 (TUBITAK 2204-B ortaokul projesi; egitim amacli)"}
KAYNAK_ONCELIK = {"iNaturalist": 0, "GBIF": 1, "Wikimedia Commons": 2}
try:
    from PIL import Image
except ImportError:      # pillow yoksa görüntü doğrulaması atlanır
    Image = None


def adaylari_oku(yol, ek_kolon=None):
    foto = pd.read_excel(yol, sheet_name="Foto_Adaylari")
    foto = foto[foto["foto_turu"] == "Canlı"]
    return {ad: g.to_dict("records") for ad, g in foto.groupby("tur_adi")}


def kalem_listesi(a):
    """Endemik (P1) ve kontrol (N1-N3) türlerini tek listede toplar."""
    t = pd.read_excel(a.endemik, sheet_name="Turler")
    t = t[t["Test_İçin_Uygun"] == "Evet"]
    ad_e = adaylari_oku(a.endemik)
    ad_k = adaylari_oku(a.kontrol)
    k = pd.read_excel(a.kontrol, sheet_name="Turler")
    kalem = []
    for _, r in t.iterrows():
        kalem.append({"tur": r["Bilimsel Adı (Latince)"], "grup": "P1", "endemik": "EVET",
                      "yayilis": r["Yayılış_Grubu"], "il_sayisi": r["İl_Sayısı"],
                      "referans": "Nuh'un Gemisi (T.C. Tarım ve Orman Bakanlığı)", "adaylar": ad_e.get(r["Bilimsel Adı (Latince)"], [])})
    for _, r in k.iterrows():
        kalem.append({"tur": r["Bilimsel Adı (Latince)"], "grup": r["Grup_Kodu"], "endemik": "HAYIR",
                      "yayilis": "", "il_sayisi": "", "referans": str(r["Endemik_Değil_Kanıtı"]),
                      "adaylar": ad_k.get(r["Bilimsel Adı (Latince)"], [])})
    return kalem


def indir_dosya(url, yol, maks_mb=6, en_az_px=300):
    try:
        r = requests.get(url, headers=HEADERS, timeout=40)
    except requests.RequestException:
        return False
    if r.status_code != 200 or not r.headers.get("content-type", "").startswith("image"):
        return False
    if len(r.content) < 5000 or len(r.content) > maks_mb * 1e6:
        return False
    with open(yol, "wb") as f:
        f.write(r.content)
    if Image is not None:
        try:
            with Image.open(yol) as im:
                im.verify()
            with Image.open(yol) as im:
                if min(im.size) < en_az_px:
                    raise ValueError("küçük")
        except Exception:
            os.remove(yol)
            return False
    return True


def asama_indir(a):
    kalem = kalem_listesi(a)
    rng = random.Random(a.tohum)
    kalem.sort(key=lambda x: (x["grup"], x["tur"]))
    rng.shuffle(kalem)
    for i, k in enumerate(kalem, 1):
        k["foto_id"] = f"F{i:04d}"             # numara türe sabit bağlanır (indirme başarısı numarayı değiştirmez)
    os.makedirs(a.foto_klasoru, exist_ok=True)
    ob_yol = os.path.join(a.calisma, "test_secim.json")
    os.makedirs(a.calisma, exist_ok=True)
    ob = json.load(open(ob_yol, encoding="utf-8")) if os.path.isfile(ob_yol) else {}
    yeni = 0
    for i, k in enumerate(kalem, 1):
        anahtar = f"{k['grup']}|{k['tur']}"
        if anahtar in ob:
            continue
        adaylar = sorted(k["adaylar"], key=lambda c: (KAYNAK_ONCELIK.get(c["kaynak"], 9), random.Random(f"{a.tohum}{c['foto_id']}").random()))
        sec = None
        for c in adaylar[:a.en_cok_deneme]:
            yol = os.path.join(a.foto_klasoru, k["foto_id"] + ".jpg")
            if str(c.get("foto_url", "")).startswith("http") and indir_dosya(c["foto_url"], yol):
                sec = c
                break
            time.sleep(a.bekle / 2)
        ob[anahtar] = {"foto_id": k["foto_id"], "aday": sec} if sec else {"foto_id": k["foto_id"], "aday": None}
        yeni += 1
        print(f"[{i}/{len(kalem)}] {k['foto_id']} {k['grup']} {k['tur']}: {'OK' if sec else 'İNDİRİLEMEDİ'}")
        if yeni % 20 == 0:
            json.dump(ob, open(ob_yol, "w", encoding="utf-8"), ensure_ascii=False)
        time.sleep(a.bekle)
    json.dump(ob, open(ob_yol, "w", encoding="utf-8"), ensure_ascii=False)
    basarisiz = [k for k, v in ob.items() if v["aday"] is None]
    print(f"\nİndirilen: {len(ob) - len(basarisiz)} | indirilemeyen (dışarıda kalan tür): {len(basarisiz)}")


def set_ata(satirlar, oran, tohum):
    """Her tabaka (grup, P1 için yayılış) içinde ayrı ayrı GELISTIRME/TEST dağıt."""
    tabaka = {}
    for s in satirlar:
        tabaka.setdefault((s["grup"], s["yayilis"]), []).append(s)
    for key, lst in sorted(tabaka.items()):
        lst.sort(key=lambda s: s["foto_id"])
        random.Random(f"{tohum}{key}").shuffle(lst)
        n_dev = round(len(lst) * oran)
        for j, s in enumerate(lst):
            s["set"] = "GELISTIRME" if j < n_dev else "TEST"


def asama_plan(a):
    kalem = {f"{k['grup']}|{k['tur']}": k for k in kalem_listesi(a)}
    ob = json.load(open(os.path.join(a.calisma, "test_secim.json"), encoding="utf-8"))
    satirlar, disarida = [], []
    for anahtar, v in ob.items():
        k = kalem.get(anahtar)
        if k is None:
            continue
        if not v["aday"]:
            disarida.append(k)
            continue
        c = v["aday"]
        satirlar.append({**k, "foto_id": v["foto_id"], "kaynak": c["kaynak"], "kayit_url": c.get("kayit_url"),
                         "lisans": c.get("lisans"), "fotografci": c.get("fotografci"), "tarih": c.get("tarih"),
                         "foto_url": c.get("foto_url"), "dogrulama": c.get("dogrulama_bilgisi")})
    set_ata(satirlar, a.gelistirme_orani, a.tohum)
    satirlar.sort(key=lambda s: s["foto_id"])

    wb = load_workbook(a.sablon)
    ws = wb["Fotograf_Kaydi"]
    son = 3 + len(satirlar) + 10
    for dv in ws.data_validations.dataValidation:          # açılır listeleri yeni satır sayısına genişlet
        sutunlar = {str(r).split(":")[0].rstrip("0123456789") for r in dv.sqref.ranges}
        dv.sqref = MultiCellRange(" ".join(f"{c}4:{c}{son}" for c in sutunlar))
    basl = {str(c.value).strip(): i for i, c in enumerate(ws[2], 1) if c.value}
    for j, s in enumerate(satirlar):
        r = 4 + j
        def yaz(ad, val):
            if ad in basl:
                ws.cell(r, basl[ad], val)
        yaz("foto_id", s["foto_id"]); yaz("dosya_adi", s["foto_id"] + ".jpg"); yaz("set", s["set"]); yaz("grup", s["grup"])
        yaz("gercek_tur", s["tur"]); yaz("gercek_endemik", s["endemik"]); yaz("referans_kaynak", s["referans"])
        yaz("kaynak_turu", "ACIK_GOZLEM"); yaz("tarih", str(s["tarih"] or "")[:10])
        yaz("not", f"{s['kaynak']} | {s['kayit_url']} | {s['lisans']} | {s['fotografci']}")

    # ek sayfalar: kaynak/atıf kaydı ve plan özeti (metrik betiği bunlara dokunmaz)
    for ad in ("Foto_Kaynak", "Plan_Ozeti"):
        if ad in wb.sheetnames:
            del wb[ad]
    fk = wb.create_sheet("Foto_Kaynak")
    fk.append(["foto_id", "tur", "grup", "yayilis_grubu", "il_sayisi", "set", "kaynak", "kayit_url", "foto_url", "lisans", "fotografci", "tarih", "dogrulama_bilgisi"])
    for s in satirlar:
        fk.append([s["foto_id"], s["tur"], s["grup"], s["yayilis"], s["il_sayisi"], s["set"], s["kaynak"], s["kayit_url"],
                   s["foto_url"], s["lisans"], s["fotografci"], str(s["tarih"] or ""), s["dogrulama"]])
    df = pd.DataFrame(satirlar)
    po = wb.create_sheet("Plan_Ozeti")
    po.append(["Grup", "Yayılış (P1)", "GELISTIRME", "TEST", "Toplam"])
    for (g, y), sub in df.groupby(["grup", "yayilis"]):
        po.append([g, y, int((sub["set"] == "GELISTIRME").sum()), int((sub["set"] == "TEST").sum()), len(sub)])
    po.append([]); po.append(["TOPLAM", "", int((df["set"] == "GELISTIRME").sum()), int((df["set"] == "TEST").sum()), len(df)])
    po.append([]); po.append(["İndirilemediği için dışarıda kalan tür", len(disarida)])
    for k in disarida[:200]:
        po.append([k["grup"], k["tur"]])
    wb.save(a.cikti)
    print(f"Yazıldı: {a.cikti} | satır: {len(satirlar)} | dışarıda kalan: {len(disarida)}")
    print(df.groupby(["grup", "set"]).size().unstack(fill_value=0).to_string())
    print(f"\nToplam sorgu = {len(satirlar)}. Günlük 500 limitle yaklaşık {-(-len(satirlar) // 500)} gün.")


def main():
    ap = argparse.ArgumentParser(description="TrEndemik test planı")
    ap.add_argument("--asama", default="hepsi", choices=["indir", "plan", "hepsi"])
    ap.add_argument("--endemik", default="endemik_veri_seti_v2.xlsx")
    ap.add_argument("--kontrol", default="kontrol_veri_seti.xlsx")
    ap.add_argument("--sablon", default="test_sablonu.xlsx")
    ap.add_argument("--cikti", default="test_sablonu_dolu.xlsx")
    ap.add_argument("--calisma", default="calisma")
    ap.add_argument("--foto-klasoru", default="fotograflar")
    ap.add_argument("--gelistirme-orani", type=float, default=0.30)
    ap.add_argument("--en-cok-deneme", type=int, default=4, help="bir tür için kaç adayı denesin")
    ap.add_argument("--tohum", type=int, default=2027)
    ap.add_argument("--bekle", type=float, default=0.3)
    a = ap.parse_args()
    if a.asama in ("indir", "hepsi"):
        asama_indir(a)
    if a.asama in ("plan", "hepsi"):
        asama_plan(a)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nDurduruldu. Aynı komutla yeniden çalıştırınca kaldığı yerden devam eder.")
